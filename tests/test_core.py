"""确定性内核测试：预算 / JSON 提取 / 黑板 / 表格比对 / 校验规则。

不依赖任何 LLM，因此可离线、可重复运行：python tests/test_core.py
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _util import make_xlsx  # noqa: E402

from agents.base import AppConfig  # noqa: E402
from agents.rules import run_rules  # noqa: E402
from core.blackboard import Blackboard  # noqa: E402
from core.budget import BudgetExceeded, TokenBudget  # noqa: E402
from core.llm import LLMResponse, MockLLM  # noqa: E402
from core.utils import extract_json  # noqa: E402
from scenarios import load_scenario  # noqa: E402
from tools import data_ops  # noqa: E402
from tools.registry import default_registry  # noqa: E402


def test_budget_accumulates_and_trips():
    b = TokenBudget(max_total_tokens=100, price_in_per_m=2.0, price_out_per_m=8.0)
    b.add("planner", 30, 30)
    assert b.total_tokens == 60
    assert b.per_agent["planner"].calls == 1
    try:
        b.add("executor", 100, 100)
        raise AssertionError("应当超出预算")
    except BudgetExceeded:
        pass
    summary = b.summary()
    assert summary["per_agent"]["planner"]["total_tokens"] == 60
    assert summary["estimated_cost_yuan"] > 0


def test_extract_json_tolerant():
    assert extract_json('{"a": 1}') == {"a": 1}
    assert extract_json('```json\n{"a": 2}\n```') == {"a": 2}
    assert extract_json('前言 {"a": 3, "b": [1,2]} 后记') == {"a": 3, "b": [1, 2]}
    # 完整对象后面跟着多余内容 → 回退到首个合法对象
    assert extract_json('{"a": 1} garbage {"b": 2}') == {"a": 1}
    # 真正缺右括号的截断无法补齐 → 返回 None（由调用方兜底）
    assert extract_json('{"a": 1, "b": {"c": 2}') is None
    assert extract_json("no json here") is None


def test_blackboard_sections_and_sources():
    bb = Blackboard("任务", "reconcile")
    bb.add_source("A", 1, "http://x", 0.9)
    bb.add_finding({"type": "not_found", "item": "B"})
    bb.add_error("executor", "boom")
    snap = bb.snapshot()
    assert snap["collected"][0]["source"] == "http://x"
    assert snap["findings"][0]["item"] == "B"
    assert snap["errors"][0]["where"] == "executor"


def test_mock_llm_scripted():
    m = MockLLM([LLMResponse(content="hi", prompt_tokens=3, completion_tokens=2)])
    r = m.chat([{"role": "user", "content": "x"}])
    assert r.content == "hi" and r.total_tokens == 5
    assert len(m.calls) == 1


def test_diff_and_duplicates():
    with tempfile.TemporaryDirectory() as d:
        left = make_xlsx(Path(d) / "left.xlsx", ["订单号", "金额", "客户"], [
            ["A001", 100, "张三"],
            ["A002", 200, "李四"],
            ["A003", 300, "王五"],
            ["A002", 200, "李四"],  # 左表重复键
        ])
        right = make_xlsx(Path(d) / "right.xlsx", ["订单号", "金额", "客户"], [
            ["A001", 100, "张三"],
            ["A002", 250, "李四"],
            ["A004", 400, "赵六"],
        ])

        diff = data_ops.diff_tables(str(left), str(right), ["订单号"], ["金额"])
        assert diff["status"] == "ok"
        assert ["A003"] in diff["only_in_left"]
        assert ["A004"] in diff["only_in_right"]
        assert any(m["key"] == ["A002"] and m["left"] == "200" and m["right"] == "250"
                   for m in diff["value_mismatches"])
        assert len(diff["duplicate_keys_left"]) == 1

        dup = data_ops.find_duplicates(str(left), key_columns=["订单号"])
        assert dup["duplicate_count"] == 1
        assert dup["duplicate_groups"][0]["rows"] == [3, 5]


def _dummy_ctx(scenario_id: str):
    return SimpleNamespace(
        registry=default_registry(),
        scenario=load_scenario(scenario_id),
        config=AppConfig(),
    )


def test_rules_reconcile():
    ctx = _dummy_ctx("reconcile")
    with tempfile.TemporaryDirectory() as d:
        left = make_xlsx(Path(d) / "left.xlsx", ["订单号", "金额"], [["A001", 100], ["A002", 200]])
        right = make_xlsx(Path(d) / "right.xlsx", ["订单号", "金额"], [["A001", 100], ["A002", 250]])
        bb = Blackboard("t", "reconcile")
        bb.set("plan", {"task_type": "reconcile", "target_files": [str(left)], "key_columns": ["订单号"]})

        # 未产出 diff → 应为 error
        findings = run_rules(["diff_consistency"], ctx, bb)
        assert any(f["severity"] == "error" for f in findings)

        # 产出 diff 后 → info
        bb.set("diff", data_ops.diff_tables(str(left), str(right), ["订单号"], ["金额"]))
        findings = run_rules(["diff_consistency"], ctx, bb)
        assert all(f["severity"] == "info" for f in findings)
        assert findings[0]["diff_count"] == 1


def test_rules_no_hallucination():
    ctx = _dummy_ctx("quote_fill")
    bb = Blackboard("t", "quote_fill")
    bb.set("collected", [{"item": "香樟", "value": 85, "source": "http://x", "confidence": 1.0}])

    # 写入与来源一致 → 通过
    bb.set("writes", [{"file": "f.xlsx", "saved": True,
                       "written": [{"row": 2, "item": "香樟", "price": 85}]}])
    findings = run_rules(["no_hallucinated_values", "values_traceable"], ctx, bb)
    assert all(f["severity"] == "info" for f in findings), findings

    # 篡改金额 → 必须被判定为 error
    bb.set("writes", [{"file": "f.xlsx", "saved": True,
                       "written": [{"row": 2, "item": "香樟", "price": 999}]}])
    findings = run_rules(["no_hallucinated_values"], ctx, bb)
    assert any(f["severity"] == "error" for f in findings)

    # 来源缺失 → values_traceable 报 error
    bb.set("collected", [{"item": "香樟", "value": 85, "source": "", "confidence": 1.0}])
    bb.set("writes", [{"file": "f.xlsx", "saved": True,
                       "written": [{"row": 2, "item": "香樟", "price": 85}]}])
    findings = run_rules(["values_traceable"], ctx, bb)
    assert any(f["severity"] == "error" for f in findings)


def test_rules_default_fill_ok():
    ctx = _dummy_ctx("quote_fill")
    bb = Blackboard("t", "quote_fill")
    bb.set("writes", [{"file": "f.xlsx", "saved": True, "written_count": 3,
                       "written": [{"row": 2, "item": "A", "price": 1}],
                       "not_found": []}])
    findings = run_rules(["no_empty_required"], ctx, bb)
    assert all(f["severity"] == "info" for f in findings), findings


def test_batch_fill_many_parallel():
    """并行写多个文件：线程池处理，结果与串行一致。"""
    import tempfile

    from excel_ops import batch_fill_many
    from openpyxl import load_workbook

    with tempfile.TemporaryDirectory() as d:
        files = []
        for i in range(6):
            files.append(str(make_xlsx(Path(d) / f"t{i}.xlsx", ["项目名称", "最新市场价"],
                                       [["香樟", None], ["桂花", None], ["银杏", None]])))
        res = batch_fill_many(files, {"香樟": 85, "桂花": 120, "银杏": 260}, workers=4)
        assert res["file_count"] == 6
        assert res["written_count"] == 18, res
        assert res["saved"] and res["failed_count"] == 0
        wb = load_workbook(files[0])
        ws = wb.active
        assert ws["B2"].value == 85 and ws["B3"].value == 120 and ws["B4"].value == 260
        wb.close()


def main() -> int:
    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_") and callable(v)]
    failed = 0
    for fn in tests:
        try:
            fn()
            print(f"  PASS  {fn.__name__}")
        except Exception as e:  # noqa: BLE001
            failed += 1
            print(f"  FAIL  {fn.__name__}: {e}")
    print(f"\n{len(tests) - failed}/{len(tests)} passed")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
