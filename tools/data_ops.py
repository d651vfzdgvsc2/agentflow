"""数据核对 / 清洗操作层（纯代码，不含 LLM 逻辑）。

这些工具为"对账"和"数据清洗"场景提供确定性的表格比对能力，
是 Verifier 判定结果的依据，也是 Agent 可以调用的原子操作。
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from excel_ops import DEFAULT_DIR, _resolve_path, _safe


def _open(path: Path, data_only: bool = True):
    return load_workbook(path, data_only=data_only, read_only=True)


def _sheet_names(path: Path) -> list[str]:
    wb = _open(path)
    names = list(wb.sheetnames)
    wb.close()
    return names


def read_records(file_name: str, sheet_name: str | None = None, base_dir: str | None = None) -> dict[str, Any]:
    """把一张 sheet 读成"表头 + 记录列表"，供比对/清洗使用。

    返回 {"file","sheet","header":[...],"records":[{列名:值}, ...], "row_numbers":[...]}
    第 1 行视为表头，数据从第 2 行开始。
    """
    path = _resolve_path(file_name, base_dir)
    if not path.exists():
        return {"error": f"文件不存在: {path}"}
    wb = _open(path)
    if sheet_name is None:
        sheet_name = wb.sheetnames[0]
    if sheet_name not in wb.sheetnames:
        wb.close()
        return {"error": f"Sheet 不存在: {sheet_name}，可用: {wb.sheetnames}"}
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return {"file": str(path), "sheet": sheet_name, "header": [], "records": [], "row_numbers": []}

    header = ["" if v is None else str(v).strip() for v in rows[0]]
    records, row_numbers = [], []
    for idx, row in enumerate(rows[1:], start=2):
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        rec = {}
        for i, name in enumerate(header):
            rec[name] = row[i] if i < len(row) else None
        records.append(rec)
        row_numbers.append(idx)
    return {
        "file": str(path),
        "sheet": sheet_name,
        "header": header,
        "records": records,
        "row_numbers": row_numbers,
    }


def _norm(v: Any) -> str:
    """用于比对的归一化：去掉首尾空白、把数字统一成去掉末尾 .0 的字符串。"""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def diff_tables(
    left_file: str,
    right_file: str,
    key_columns: list[str],
    compare_columns: list[str] | None = None,
    left_sheet: str | None = None,
    right_sheet: str | None = None,
    base_dir: str | None = None,
) -> dict[str, Any]:
    """按关键列比对两张表，返回差异清单（缺失行 / 仅单边存在 / 字段不一致 / 重复键）。

    说明：只做确定性比对，不做任何"猜测判定"。
    """
    left = read_records(left_file, left_sheet, base_dir)
    if "error" in left:
        return {"error": f"左表读取失败: {left['error']}"}
    right = read_records(right_file, right_sheet, base_dir)
    if "error" in right:
        return {"error": f"右表读取失败: {right['error']}"}

    compare_columns = compare_columns or []

    def _key(rec: dict) -> tuple:
        return tuple(_norm(rec.get(k)) for k in key_columns)

    def _index(records: list[dict]) -> dict[tuple, list[dict]]:
        idx: dict[tuple, list[dict]] = {}
        for r in records:
            idx.setdefault(_key(r), []).append(r)
        return idx

    li, ri = _index(left["records"]), _index(right["records"])

    only_left = [k for k in li if k not in ri]
    only_right = [k for k in ri if k not in li]
    dup_left = [{"key": list(k), "count": len(v)} for k, v in li.items() if len(v) > 1]
    dup_right = [{"key": list(k), "count": len(v)} for k, v in ri.items() if len(v) > 1]

    mismatches = []
    for k in li:
        if k not in ri:
            continue
        lrec, rrec = li[k][0], ri[k][0]
        for col in compare_columns:
            lv, rv = _norm(lrec.get(col)), _norm(rrec.get(col))
            if lv != rv:
                mismatches.append({"key": list(k), "column": col, "left": lv, "right": rv,
                                   "left_row": lrec, "right_row": rrec})

    return {
        "status": "ok",
        "key_columns": key_columns,
        "compare_columns": compare_columns,
        "header": left["header"],
        "left": {"file": left["file"], "sheet": left["sheet"], "rows": len(left["records"])},
        "right": {"file": right["file"], "sheet": right["sheet"], "rows": len(right["records"])},
        "only_in_left": [list(k) for k in only_left],
        "only_in_left_rows": [li[k][0] for k in only_left],
        "only_in_right": [list(k) for k in only_right],
        "only_in_right_rows": [ri[k][0] for k in only_right],
        "duplicate_keys_left": dup_left,
        "duplicate_keys_right": dup_right,
        "value_mismatches": mismatches,
        "diff_count": len(only_left) + len(only_right) + len(mismatches) + len(dup_left) + len(dup_right),
    }


def find_duplicates(
    file_name: str,
    sheet_name: str | None = None,
    key_columns: list[str] | None = None,
    base_dir: str | None = None,
) -> dict[str, Any]:
    """在单张表里按关键列找重复记录。"""
    data = read_records(file_name, sheet_name, base_dir)
    if "error" in data:
        return data
    key_columns = key_columns or data["header"][:1]

    groups: dict[tuple, list[int]] = {}
    for rec, rn in zip(data["records"], data["row_numbers"]):
        k = tuple(_norm(rec.get(c)) for c in key_columns)
        groups.setdefault(k, []).append(rn)

    dupes = [{"key": list(k), "rows": rows} for k, rows in groups.items() if len(rows) > 1]
    return {
        "status": "ok",
        "file": data["file"],
        "sheet": data["sheet"],
        "key_columns": key_columns,
        "duplicate_groups": dupes,
        "duplicate_count": len(dupes),
    }


def fill_missing(
    file_name: str,
    sheet_name: str,
    column: str,
    source_map: dict[str, Any],
    key_column: str,
    base_dir: str | None = None,
) -> dict[str, Any]:
    """把 source_map 中"键→值"按 key_column 填进目标 column 的空单元格（只填空、不覆盖已有值）。

    用于"数据清洗补全"场景：清洗的补全动作必须可追溯，因此只补空值。
    """
    path = _resolve_path(file_name, base_dir)
    if not path.exists():
        return {"error": f"文件不存在: {path}"}
    wb = load_workbook(path, data_only=False)
    if sheet_name not in wb.sheetnames:
        wb.close()
        return {"error": f"Sheet 不存在: {sheet_name}"}
    ws = wb[sheet_name]

    header = {}
    from openpyxl.utils.cell import get_column_letter

    for i, cell in enumerate(ws[1], start=1):
        if cell.value is not None:
            header[str(cell.value).strip()] = get_column_letter(i)
    if key_column not in header or column not in header:
        wb.close()
        return {"error": f"列不存在: key={key_column}, target={column}，可用表头={list(header)}"}

    key_col, target_col = header[key_column], header[column]
    from excel_ops import _ensure_backup

    _ensure_backup(path)
    filled, skipped = [], []
    for row in range(2, ws.max_row + 1):
        k = ws[f"{key_col}{row}"].value
        if k is None:
            continue
        k = str(k).strip()
        cur = ws[f"{target_col}{row}"].value
        if cur is not None and str(cur).strip() != "":
            skipped.append({"row": row, "item": k, "reason": "已有值"})
            continue
        if k in source_map:
            ws[f"{target_col}{row}"] = _safe(source_map[k])
            filled.append({"row": row, "item": k, "value": source_map[k]})
    wb.save(path)
    wb.close()
    return {
        "file": str(path),
        "sheet": sheet_name,
        "filled": filled,
        "filled_count": len(filled),
        "skipped": skipped,
        "skipped_count": len(skipped),
        "saved": True,
    }


def merge_complete(
    left_file: str,
    right_file: str,
    key_columns: list[str] | None = None,
    compare_columns: list[str] | None = None,
    left_sheet: str | None = None,
    right_sheet: str | None = None,
    base_dir: str | None = None,
) -> dict[str, Any]:
    """按关键列把两张表合并成一张"完整表"（以左表=基准为准）。

    规则（可追溯、不改原表）：
    - 左表为基准，行顺序保持不变；
    - 左表**空**单元格用右表同键的值补全（只补空，绝不覆盖已有值）；
    - 右表独有的键，追加到结果末尾，并在「来源」列标注；
    - 结果写到 data/output/ 下的**新文件**，不动任何原始表。
    """
    from openpyxl import Workbook

    left = read_records(left_file, left_sheet, base_dir)
    if "error" in left:
        return {"error": f"左表读取失败: {left['error']}"}
    right = read_records(right_file, right_sheet, base_dir)
    if "error" in right:
        return {"error": f"右表读取失败: {right['error']}"}

    key_columns = key_columns or [left["header"][0]]

    # 表头：先左后右，合并去重，末尾加「来源」
    header = list(left["header"])
    for h in right["header"]:
        if h not in header:
            header.append(h)
    src_col = "来源"
    if src_col in header:
        src_col = "来源(合并)"
    header_out = header + [src_col]

    def _key(rec: dict) -> tuple:
        return tuple(_norm(rec.get(k)) for k in key_columns)

    ridx: dict[tuple, dict] = {}
    for r in right["records"]:
        ridx.setdefault(_key(r), r)

    filled: list[dict] = []
    out_rows: list[dict] = []
    left_keys = set()

    for lr in left["records"]:
        k = _key(lr)
        left_keys.add(k)
        row = {h: lr.get(h) for h in header}
        rr = ridx.get(k)
        if rr:
            for h in header:
                lv, rv = row.get(h), rr.get(h)
                if (lv is None or str(lv).strip() == "") and rv is not None and str(rv).strip() != "":
                    row[h] = rv
                    filled.append({"key": list(k), "column": h, "value": rv})
        row[src_col] = "基准表"
        out_rows.append(row)

    added: list[list] = []
    for rr in right["records"]:
        k = _key(rr)
        if k in left_keys:
            continue
        row = {h: rr.get(h) for h in header}
        row[src_col] = "仅待核对表"
        out_rows.append(row)
        added.append(list(k))

    # 写新文件（绝不改原表）
    import time

    out_dir = DEFAULT_DIR / "output"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"合并补齐_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "合并结果"
    ws.append(header_out)
    for row in out_rows:
        ws.append([_safe(row.get(h)) for h in header_out])
    wb.save(out_path)
    wb.close()

    return {
        "status": "ok",
        "output": str(out_path),
        "output_name": out_path.name,
        "key_columns": key_columns,
        "header": header_out,
        "rows": len(out_rows),
        "base_rows": len(left["records"]),
        "added_from_right": len(added),
        "added_keys": added[:200],
        "filled_count": len(filled),
        "filled": filled[:200],
    }


def list_sheets(file_name: str, base_dir: str | None = None) -> dict[str, Any]:
    """列出文件里的所有 sheet 名。"""
    path = _resolve_path(file_name, base_dir)
    if not path.exists():
        return {"error": f"文件不存在: {path}"}
    return {"file": str(path), "sheets": _sheet_names(path)}


# 兼容：DEFAULT_DIR 从 excel_ops 再导出，方便调用方统一入口
__all__ = ["read_records", "diff_tables", "find_duplicates", "fill_missing",
           "merge_complete", "list_sheets", "DEFAULT_DIR"]
